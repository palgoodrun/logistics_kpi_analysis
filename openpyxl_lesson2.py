from openpyxl import Workbook, load_workbook
from pathlib import Path

output_path = Path("output") / "openpyxl_lesson1.xlsx"

wb = load_workbook(output_path)

print(wb.worksheets)


summary_ws = wb["Summary"]

print(summary_ws.max_row, summary_ws.max_column)


print(summary_ws["B4"].value)


rows = summary_ws["A3:C5"]

for row in rows:
    for cell in row:
        print(cell.value)


rows = summary_ws.iter_rows(
    min_row=3,
    max_row=summary_ws.max_row,
    min_col=1,
    max_col=3,
    values_only=True,
)

for row in rows:
    print(row)


detail_ws = wb["Detail"]

print(detail_ws["A1"].value)
