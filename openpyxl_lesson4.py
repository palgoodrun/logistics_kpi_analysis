import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.cell.cell import Cell
from pathlib import Path

df = pd.DataFrame(
    [
        ["Carrier A", 1200, 150000],
        ["Carrier B", 800, 110000],
        ["Carrier C", 1500, 180000],
        ["Carrier D", 600, 90000],
    ],
    columns=["配送会社", "配送数量", "運賃"],
)

wb = Workbook()
ws = wb.active
ws.title = "Summary"

ws["A1"] = "物流KPI分析レポート"
ws["A2"] = "配送会社別集計"

for row_index, row in enumerate(
    dataframe_to_rows(df, index=False, header=True), start=4
):
    for col_index, value in enumerate(row, start=1):
        ws.cell(row=row_index, column=col_index, value=value)

header_font = Font(
    bold=True,
    color="FFFFFF",
)
header_fill = PatternFill(
    "solid",
    fgColor="24466B",
)
header_alignment = Alignment(horizontal="center")

for row in ws["A4:C4"]:
    for cell in row:
        cell: Cell
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

for row in ws.iter_rows(min_row=5, max_row=ws.max_row, min_col=2, max_col=2):
    for cell in row:
        cell: Cell
        cell.number_format = "#,##0"

for row in ws.iter_rows(min_row=5, max_row=ws.max_row, min_col=3, max_col=3):
    for cell in row:
        cell: Cell
        cell.number_format = '"¥"#,##0'

ws.column_dimensions["A"].width = 25
ws.column_dimensions["B"].width = 15
ws.column_dimensions["C"].width = 20

ws.row_dimensions[1].height = 30
ws.freeze_panes = "A5"

output_dir = "output"
output_file = "openpyxl_lesson4.xlsx"
output_path = Path(output_dir) / output_file

wb.save(output_path)

wb = load_workbook(output_path)
ws = wb.active

print(ws.title)
print(ws["A4"].value)
print(ws["A7"].value)
print(ws["B7"].value)
print(ws["C7"].number_format)
print(ws.freeze_panes)
