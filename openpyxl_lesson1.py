from openpyxl import Workbook, load_workbook
from pathlib import Path

wb = Workbook()
summary_ws = wb.active
summary_ws.title = "Summary"


summary_ws["A1"] = "物流KPI分析レポート"
summary_ws.cell(row=3, column=1, value="配送会社")
summary_ws.cell(row=3, column=2, value="配送数量")
summary_ws.cell(row=3, column=3, value="運賃")
summary_ws.cell(row=4, column=1, value="Carrier A")
summary_ws.cell(row=4, column=2, value=1200)
summary_ws.cell(row=4, column=3, value=150000)
summary_ws.cell(row=5, column=1, value="Carrier B")
summary_ws.cell(row=5, column=2, value=800)
summary_ws.cell(row=5, column=3, value=110000)


detail_ws = wb.create_sheet("Detail")

detail_ws["A1"] = "Detail Data"


output_dir = Path("output")
output_dir.mkdir(exist_ok=True)
output_path = output_dir / "openpyxl_lesson1.xlsx"


wb.save(output_path)


loaded_wb = load_workbook(output_path)
loaded_summary_ws = loaded_wb["Summary"]
loaded_detail = loaded_wb["Detail"]

print(loaded_summary_ws["B4"].value)
print(loaded_detail["A1"].value)
