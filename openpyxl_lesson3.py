from pathlib import Path
from openpyxl.cell.cell import Cell
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

output_dir = "output"

input_file = "openpyxl_lesson1.xlsx"
input_path = Path(output_dir) / input_file

wb = load_workbook(input_path)
ws = wb["Summary"]


ws["A1"].font = Font(
    size=16,
    bold=True,
    color="0000FF",
)


rows = ws["A3:C3"]

for row in rows:
    for cell in row:
        cell: Cell
        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="24466B",
        )
        cell.border = Border(
            bottom=Side(
                style="thin",
                color="000000",
            )
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )


rows = ws["B4:B5"]

for row in rows:
    for cell in row:
        cell: Cell
        cell.number_format = "#,##0"


rows = ws["C4:C5"]

for row in rows:
    for cell in row:
        cell: Cell
        cell.number_format = '"¥"#,##0'


ws.column_dimensions["A"].width = 25
ws.column_dimensions["B"].width = 15
ws.column_dimensions["C"].width = 20

ws.row_dimensions[1].height = 30
ws.row_dimensions[3].height = 25

ws.freeze_panes = "B4"

output_file = "openpyxl_lesson3.xlsx"
output_path = Path(output_dir) / output_file

wb.save(output_path)

wb = load_workbook(output_path)
ws = wb["Summary"]

print(ws["A1"].value)
print(ws["A1"].font.bold)
print(ws["B4"].value)
print(ws["B4"].number_format)
print(ws["C4"].number_format)
print(ws.freeze_panes)
