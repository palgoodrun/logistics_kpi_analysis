# ==================================================
# import
# ==================================================
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Font, PatternFill, Alignment
from pathlib import Path


# ==================================================
# write_dataframe()
# ==================================================
def write_dataframe(
    ws: Worksheet,
    df: pd.DataFrame,
    start_row: int = 4,
    start_col: int = 1,
) -> None:

    for row_index, row in enumerate(
        dataframe_to_rows(df, index=False, header=True), start=start_row
    ):
        for col_index, value in enumerate(row, start=start_col):
            ws.cell(
                row=row_index,
                column=col_index,
                value=value,
            )


# ==================================================
# apply_header_style()
# ==================================================
def apply_header_style(ws: Worksheet, header_row: str = 4) -> None:

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="24466B")
    header_alignment = Alignment(horizontal="center")

    for row in ws.iter_rows(
        min_row=header_row,
        max_row=header_row,
        min_col=1,
        max_col=ws.max_column,
    ):
        for cell in row:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment


# ==================================================
# setting_worksheets()
# ==================================================
def setting_worksheets(
    data: dict,
    wb: Workbook,
) -> None:

    for sheet_name in data:
        wb[sheet_name]["A1"] = f"{sheet_name} 集計レポート"
        wb[sheet_name]["A2"] = "物流KPI分析"

        wb[sheet_name].column_dimensions["A"].width = 25
        wb[sheet_name].column_dimensions["B"].width = 15
        wb[sheet_name].column_dimensions["C"].width = 20
        wb[sheet_name].row_dimensions[1].height = 30
        wb[sheet_name].freeze_panes = "A5"


# ==================================================
# main()
# ==================================================
def main():

    # ==================================================
    # dfs
    # ==================================================
    df_carrier = pd.DataFrame(
        {
            "配送会社": ["Carrier A", "Carrier B", "Carrier C"],
            "配送数量": [1200, 800, 1500],
            "運賃": [150000, 110000, 180000],
        }
    )

    df_area = pd.DataFrame(
        {
            "エリア": ["関東", "関西", "九州"],
            "配送数量": [1100, 900, 500],
            "配送比率": [0.44, 0.36, 0.20],
        }
    )

    df_product = pd.DataFrame(
        {
            "商品": ["水", "サーバー", "フィルター"],
            "配送数量": [1500, 800, 700],
            "運賃": [200000, 140000, 100000],
        }
    )

    df_month = pd.DataFrame(
        {
            "month": ["1月", "2月"],
            "配送数量": [1000, 1100],
            "運賃": [120000, 130000],
        }
    )

    # ==================================================
    # report_data:dict
    # ==================================================
    report_data = {
        "Carrier": df_carrier,
        "Area": df_area,
        "Product": df_product,
        "Month": df_month,
    }

    # ==================================================
    # wb / ws
    # ==================================================
    wb = Workbook()
    wb.remove(wb.active)

    for sheet_name, df in report_data.items():
        ws = wb.create_sheet(sheet_name)

        write_dataframe(ws, df)

        apply_header_style(ws)

    setting_worksheets(report_data, wb)

    output_dir = "output"
    output_file = "openpyxl_lesson5.xlsx"
    output_path = Path(output_dir) / output_file

    for row in wb["Carrier"].iter_rows(
        min_row=5, max_row=wb["Carrier"].max_row, min_col=2, max_col=2
    ):
        for cell in row:
            cell.number_format = "#,##0"

    for row in wb["Carrier"].iter_rows(
        min_row=5, max_row=wb["Carrier"].max_row, min_col=3, max_col=3
    ):
        for cell in row:
            cell.number_format = '"¥"#,##0'

    for row in wb["Area"].iter_rows(
        min_row=5, max_row=wb["Area"].max_row, min_col=2, max_col=2
    ):
        for cell in row:
            cell.number_format = "#,##0"

    for row in wb["Area"].iter_rows(
        min_row=5, max_row=wb["Area"].max_row, min_col=3, max_col=3
    ):
        for cell in row:
            cell.number_format = "0.0%"

    for row in wb["Product"].iter_rows(
        min_row=5, max_row=wb["Product"].max_row, min_col=2, max_col=2
    ):
        for cell in row:
            cell.number_format = "#,##0"

    for row in wb["Product"].iter_rows(
        min_row=5, max_row=wb["Product"].max_row, min_col=3, max_col=3
    ):
        for cell in row:
            cell.number_format = '"¥"#,##0'

    for row in wb["Month"].iter_rows(
        min_row=5, max_row=wb["Month"].max_row, min_col=2, max_col=2
    ):
        for cell in row:
            cell.number_format = "#,##0"

    for row in wb["Month"].iter_rows(
        min_row=5, max_row=wb["Month"].max_row, min_col=3, max_col=3
    ):
        for cell in row:
            cell.number_format = '"¥"#,##0'

    wb.save(output_path)

    wb = load_workbook(output_path)
    ws_carrier = wb["Carrier"]
    ws_area = wb["Area"]
    ws_product = wb["Product"]

    print(wb.sheetnames)
    print(ws_carrier.max_row, ws_carrier.max_column)
    print(ws_area.max_row, ws_area.max_column)
    print(ws_product.max_row, ws_product.max_column)
    print(ws_carrier["A4"].value)
    print(ws_area["C5"].value, ws_area["C5"].number_format)
    print(ws_product["A7"].value)
    print(ws_carrier.freeze_panes, ws_area.freeze_panes, ws_product.freeze_panes)


# ==================================================
# entrypoint
# ==================================================
if __name__ == "__main__":
    main()
